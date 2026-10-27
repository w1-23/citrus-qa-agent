"""Citrus research workflow tools and session-scoped downloadable artifacts.

Uploaded tabular data is parsed in-process. Model-generated code is never run.
The model interprets field semantics; deterministic checks report findings
without modifying the uploaded source.
"""
from __future__ import annotations

import csv
import base64
import hashlib
import html
import io
import json
import logging
import math
import re
import sqlite3
import tarfile
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile
from difflib import SequenceMatcher
from itertools import islice
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

import numpy as np

from langchain_core.tools import tool

from src.config import PROJECT_ROOT
from src.core.tracing import get_session_id, get_user_id
from src.research.context import get_request_tables

logger = logging.getLogger(__name__)
_ARTIFACT_TTL_SEC = 24 * 60 * 60
_MAX_TABLE_ROWS = 100_000
_MAX_TABLE_COLUMNS = 200
_MAX_TABLE_CELLS = 2_000_000
_MAX_EXTRACTED_TABLES = 500
_MAX_SAMPLES = 8


def _db_path() -> Path:
    path = PROJECT_ROOT / "state" / "research_artifacts.sqlite3"
    path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(str(path)) as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS artifacts (
            artifact_id TEXT PRIMARY KEY, session_id TEXT NOT NULL, filename TEXT NOT NULL,
            mime TEXT NOT NULL, created_at REAL NOT NULL, path TEXT NOT NULL)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS research_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, session_id TEXT NOT NULL,
            owner_user_id INTEGER, project_name TEXT NOT NULL,
            note TEXT NOT NULL, created_at REAL NOT NULL)""")
        columns = {row[1] for row in conn.execute("PRAGMA table_info(research_notes)")}
        if "owner_user_id" not in columns:
            conn.execute("ALTER TABLE research_notes ADD COLUMN owner_user_id INTEGER")
            sessions_path = PROJECT_ROOT / "state" / "sessions.db"
            if sessions_path.exists():
                try:
                    conn.execute("ATTACH DATABASE ? AS session_store", (str(sessions_path),))
                    session_columns = {row[1] for row in conn.execute(
                        "PRAGMA session_store.table_info(sessions)")}
                    if "owner_user_id" in session_columns:
                        conn.execute("""UPDATE research_notes
                            SET owner_user_id=(SELECT s.owner_user_id FROM session_store.sessions s
                                               WHERE s.session_id=research_notes.session_id)
                            WHERE owner_user_id IS NULL AND EXISTS (
                                SELECT 1 FROM session_store.sessions s
                                WHERE s.session_id=research_notes.session_id
                                  AND s.owner_user_id IS NOT NULL)""")
                        conn.commit()
                    conn.execute("DETACH DATABASE session_store")
                except sqlite3.Error:
                    logger.warning("Could not migrate legacy research-note ownership; unverified rows remain private")
                    try:
                        conn.execute("DETACH DATABASE session_store")
                    except sqlite3.Error:
                        pass
        conn.execute("CREATE INDEX IF NOT EXISTS idx_research_notes_owner_project "
                     "ON research_notes(owner_user_id,project_name,id)")
    return path


def _safe_name(value: str, fallback: str = "research-output") -> str:
    value = re.sub(r"[^\w\-. ]+", "_", str(value or ""), flags=re.UNICODE).strip(" .")
    value = re.sub(r"\s+", "_", value)[:80]
    return value or fallback


_SVG_TAGS = {"svg", "g", "path", "rect", "circle", "ellipse", "line", "polyline",
             "polygon", "text", "tspan", "title", "desc", "defs", "linearGradient",
             "radialGradient", "stop", "clipPath", "marker", "use", "symbol", "filter",
             "feGaussianBlur", "feOffset", "feFlood", "feComposite", "feMerge", "feMergeNode",
             "feColorMatrix", "feBlend", "feComponentTransfer", "feFuncA", "feFuncR", "feFuncG", "feFuncB"}
_SVG_ATTRS = {"id", "viewBox", "preserveAspectRatio", "width", "height", "x", "y", "x1", "y1",
              "x2", "y2", "cx", "cy", "r", "rx", "ry", "d", "points", "fill", "stroke",
              "stroke-width", "stroke-linecap", "stroke-linejoin", "stroke-dasharray", "stroke-dashoffset",
              "fill-opacity", "stroke-opacity", "opacity", "transform", "offset", "font-family", "font-size",
              "font-weight", "text-anchor", "dominant-baseline", "gradientUnits", "gradientTransform",
              "stop-color", "stop-opacity", "clipPathUnits", "markerWidth", "markerHeight", "markerUnits",
              "marker-start", "marker-mid", "marker-end", "clip-path", "filter", "role", "aria-label", "class",
              "style", "href"}
_SVG_STYLE = {"fill", "stroke", "stroke-width", "stroke-linecap", "stroke-linejoin", "stroke-dasharray",
              "stroke-dashoffset", "fill-opacity", "stroke-opacity", "opacity", "font-family", "font-size",
              "font-weight", "text-anchor", "dominant-baseline", "stop-color", "stop-opacity", "clip-path",
              "filter", "marker-start", "marker-mid", "marker-end"}


def _sanitize_svg_for_export(source: str) -> str:
    raw = str(source or "").strip()
    if not raw or len(raw) > 100_000 or re.search(r"<!DOCTYPE|<!ENTITY", raw, re.I):
        return ""
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return ""
    if root.tag.split("}")[-1].lower() != "svg":
        return ""

    def clean(node):
        tag = node.tag.split("}")[-1]
        if tag not in _SVG_TAGS:
            return None
        out = ET.Element(tag)
        for key, raw_value in node.attrib.items():
            name = key.split("}")[-1]
            value = str(raw_value)
            if name not in _SVG_ATTRS or len(value) > 4000 or re.search(r"[<>\x00-\x1f]", value):
                continue
            if name == "href" and not re.fullmatch(r"#[A-Za-z][\w:.-]{0,100}", value):
                continue
            if name == "id" and not re.fullmatch(r"[A-Za-z][\w:.-]{0,100}", value):
                continue
            if name == "class" and not re.fullmatch(r"[A-Za-z0-9 _-]{1,200}", value):
                continue
            if re.search(r"url\s*\(", value, re.I) and not re.fullmatch(
                    r"url\(\s*#[A-Za-z][\w:.-]{0,100}\s*\)", value, re.I):
                continue
            if re.search(r"javascript:|data:|expression|@import", value, re.I):
                continue
            if name == "style":
                declarations = []
                for item in value.split(";"):
                    if ":" not in item:
                        continue
                    prop, val = (part.strip() for part in item.split(":", 1))
                    if prop.lower() in _SVG_STYLE and val and not re.search(r"[{}]", val):
                        declarations.append(f"{prop.lower()}:{val}")
                value = ";".join(declarations)
                if not value:
                    continue
            out.set(name, value)
        out.text = node.text
        for child in list(node):
            safe_child = clean(child)
            if safe_child is not None:
                out.append(safe_child)
        return out

    safe = clean(root)
    if safe is None or not (safe.get("viewBox") or safe.get("width")):
        return ""
    safe.set("xmlns", "http://www.w3.org/2000/svg")
    return ET.tostring(safe, encoding="unicode", short_empty_elements=True)


def _markdown_to_safe_html(content: str) -> str:
    """Render the common Markdown used by exports; only sanitized inline SVG is raw HTML."""
    svg_blocks: list[tuple[str, str]] = []

    def stash_svg(raw_svg: str):
        token = f"CITRUS_EXPORT_SVG_{len(svg_blocks)}_TOKEN"
        svg_blocks.append((token, _sanitize_svg_for_export(raw_svg)))
        return f"\n\n{token}\n\n"

    source = re.sub(r"```svg\s*\n([\s\S]*?)```",
                    lambda match: stash_svg(match.group(1)), str(content or ""), flags=re.I)
    source = re.sub(r"<svg\b[\s\S]*?</svg\s*>",
                    lambda match: stash_svg(match.group(0)), source, flags=re.I)

    def inline(value: str) -> str:
        value = html.escape(value, quote=True)
        value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
        value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
        value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", value)
        value = re.sub(r"\[([^\]]+)\]\((https?://[^\s)]+)\)",
                       lambda m: f'<a href="{m.group(2)}">{m.group(1)}</a>', value)
        return value

    lines = source.splitlines()
    output: list[str] = []
    paragraph: list[str] = []
    list_tag = ""
    i = 0

    def flush_paragraph():
        if paragraph:
            output.append("<p>" + " ".join(inline(line.strip()) for line in paragraph) + "</p>")
            paragraph.clear()

    def close_list():
        nonlocal list_tag
        if list_tag:
            output.append(f"</{list_tag}>")
            list_tag = ""

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_paragraph(); close_list(); i += 1; continue
        if stripped.startswith("CITRUS_EXPORT_SVG_") and stripped.endswith("_TOKEN"):
            flush_paragraph(); close_list(); output.append(stripped); i += 1; continue
        if stripped.startswith("```"):
            flush_paragraph(); close_list(); code = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            output.append("<pre><code>" + html.escape("\n".join(code)) + "</code></pre>")
            i += 1; continue
        if i + 1 < len(lines) and "|" in line and re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*", lines[i + 1]):
            flush_paragraph(); close_list()
            headers = [cell.strip() for cell in stripped.strip("|").split("|")]
            output.append("<table><thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in headers) + "</tr></thead><tbody>")
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                output.append("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in cells) + "</tr>")
                i += 1
            output.append("</tbody></table>"); continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush_paragraph(); close_list(); level = len(heading.group(1))
            output.append(f"<h{level}>{inline(heading.group(2))}</h{level}>"); i += 1; continue
        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if bullet or numbered:
            flush_paragraph(); wanted = "ul" if bullet else "ol"
            if list_tag != wanted:
                close_list(); list_tag = wanted; output.append(f"<{wanted}>")
            output.append("<li>" + inline((bullet or numbered).group(1)) + "</li>")
            i += 1; continue
        close_list(); paragraph.append(line); i += 1
    flush_paragraph(); close_list()
    rendered = "\n".join(output)
    for token, svg in svg_blocks:
        rendered = rendered.replace(token, svg or '<p class="muted">SVG 未通过安全检查，未嵌入。</p>')
    return rendered


def _safe_spreadsheet_value(value: Any) -> Any:
    """Prevent formula injection when exported files are opened in office apps."""
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@", "\t", "\r"}:
        return "'" + value
    return value


def _normalize_headers(raw_rows: list[list[Any]], header_rows: list[int], width: int) -> list[str]:
    """Combine one or more 1-based header rows, including merged/blank cells."""
    selected = []
    for row_no in sorted(set(header_rows)):
        row = raw_rows[row_no - 1] if 0 < row_no <= len(raw_rows) else []
        values = [str(v).strip() if v is not None else "" for v in row]
        values = (values + [""] * width)[:width]
        if sum(bool(v) for v in values) <= 1:  # title/date note rather than a field header
            continue
        selected.append(values)
    if not selected:
        selected = [[str(v).strip() if v is not None else ""
                     for v in (raw_rows[header_rows[-1] - 1] if header_rows else [])]]
    labels = []
    for column in range(width):
        parts = []
        carry = ""
        for row in selected:
            value = row[column] if column < len(row) else ""
            if value:
                carry = value
            if carry and (not parts or carry != parts[-1]):
                parts.append(carry)
        label = " / ".join(parts) or f"未命名列{column + 1}"
        labels.append(label)
    seen: dict[str, int] = {}
    for i, label in enumerate(labels):
        seen[label] = seen.get(label, 0) + 1
        if seen[label] > 1:
            labels[i] = f"{label} ({seen[label]})"
    return labels


def _save_artifact(filename: str, content: bytes, mime: str) -> str:
    session_id = str(get_session_id() or "")
    if not re.fullmatch(r"[A-Za-z0-9_-]{8,128}", session_id):
        raise ValueError("当前请求没有有效会话，无法保存下载文件")
    artifact_id = uuid.uuid4().hex
    root = (PROJECT_ROOT / "state" / "artifacts" / session_id).resolve()
    root.mkdir(parents=True, exist_ok=True)
    path = (root / f"{artifact_id}_{_safe_name(filename)}").resolve()
    if not path.is_relative_to(root):
        raise ValueError("无效的导出文件名")
    path.write_bytes(content)
    now = time.time()
    db = _db_path()
    with sqlite3.connect(str(db)) as conn:
        expired = conn.execute("SELECT path FROM artifacts WHERE created_at < ?",
                               (now - _ARTIFACT_TTL_SEC,)).fetchall()
        conn.execute("DELETE FROM artifacts WHERE created_at < ?", (now - _ARTIFACT_TTL_SEC,))
        conn.execute("INSERT INTO artifacts VALUES (?, ?, ?, ?, ?, ?)",
                     (artifact_id, session_id, path.name, mime, now, str(path)))
    for (old_path,) in expired:
        try:
            Path(old_path).unlink(missing_ok=True)
        except OSError:
            logger.info("Could not remove expired research artifact %s", old_path)
    return f"/api/v2/session/{session_id}/artifacts/{artifact_id}"


def resolve_artifact(session_id: str, artifact_id: str) -> dict[str, Any] | None:
    if not re.fullmatch(r"[A-Za-z0-9_-]{8,128}", str(session_id or "")) or not re.fullmatch(
            r"[a-f0-9]{32}", str(artifact_id or "")):
        return None
    now = time.time()
    with sqlite3.connect(str(_db_path())) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT * FROM artifacts WHERE artifact_id=? AND session_id=?",
                           (artifact_id, session_id)).fetchone()
        if not row:
            return None
        if now - float(row["created_at"]) > _ARTIFACT_TTL_SEC:
            conn.execute("DELETE FROM artifacts WHERE artifact_id=?", (artifact_id,))
            try:
                Path(row["path"]).unlink(missing_ok=True)
            except OSError:
                pass
            return None
    root = (PROJECT_ROOT / "state" / "artifacts" / session_id).resolve()
    path = Path(row["path"]).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        return None
    return {"path": path, "filename": row["filename"], "mime": row["mime"]}


def cleanup_expired_artifacts() -> int:
    """Remove expired generated downloads and their index rows."""
    now = time.time()
    db = _db_path()
    with sqlite3.connect(str(db)) as conn:
        rows = conn.execute("SELECT artifact_id,session_id,path FROM artifacts WHERE created_at < ?",
                            (now - _ARTIFACT_TTL_SEC,)).fetchall()
        conn.execute("DELETE FROM artifacts WHERE created_at < ?", (now - _ARTIFACT_TTL_SEC,))
    removed = 0
    for artifact_id, session_id, raw_path in rows:
        try:
            root = (PROJECT_ROOT / "state" / "artifacts" / session_id).resolve()
            path = Path(raw_path).resolve()
            if path.is_relative_to(root) and path.is_file():
                path.unlink()
                removed += 1
        except (OSError, ValueError):
            logger.warning("Could not clean expired artifact id=%s", artifact_id)
    return removed


def _attachment(file_index: int) -> dict[str, Any]:
    tables = get_request_tables()
    if not tables:
        raise ValueError("本轮没有可审查的表格或论文附件")
    if file_index < 0 or file_index >= len(tables):
        raise ValueError(f"审查文件序号无效；本轮共有 {len(tables)} 个审查附件")
    return tables[file_index]


def _decode_csv(data: bytes, header_row: int = 1, header_rows: list[int] | None = None):
    text = None
    encoding = ""
    for candidate in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            text = data.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别（尝试 UTF-8 与常见中文编码均失败）")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    header_rows = sorted(set(header_rows or [header_row]))
    if not header_rows or any(row < 1 or row > 20 for row in header_rows):
        raise ValueError("header_row 必须在 1 到 20 之间")
    row_limit = _MAX_TABLE_ROWS + max(header_rows) + 1
    raw_rows = []
    raw_cells = 0
    truncated = False
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        if len(raw_rows) >= row_limit:
            truncated = True
            break
        if len(row) > _MAX_TABLE_COLUMNS:
            raise ValueError(f"列数超过限制 {_MAX_TABLE_COLUMNS}")
        if raw_cells + len(row) > _MAX_TABLE_CELLS:
            truncated = True
            break
        raw_rows.append(row)
        raw_cells += len(row)
    raw_preview = raw_rows[:8]
    if not raw_rows:
        return _pd().DataFrame(), encoding, raw_preview, header_rows, False
    header = _normalize_headers(raw_rows, header_rows, max(len(row) for row in raw_rows))
    body_start = max(header_rows)
    width = len(header)
    if width > _MAX_TABLE_COLUMNS:
        raise ValueError(f"列数 {width} 超过限制 {_MAX_TABLE_COLUMNS}")
    data_row_limit = min(_MAX_TABLE_ROWS, _MAX_TABLE_CELLS // max(1, width))
    body = [(row + [""] * width)[:width]
            for row in raw_rows[body_start:body_start + data_row_limit]]
    frame = _pd().DataFrame(body, columns=header)
    frame = frame.replace(r"^\s*$", _pd().NA, regex=True)
    # Keep literal CSV values (including zero-padded IDs) intact. Numeric
    # calculations use temporary coercions and the raw source is never rewritten.
    truncated = truncated or len(raw_rows) > body_start + data_row_limit
    return frame, encoding, raw_preview, header_rows, truncated


def _pd():
    import pandas as pd
    return pd


def _bound_raw_table(rows: list[list[Any]], header_rows: list[int],
                     cell_budget: int, label: str) -> tuple[list[list[Any]], bool, int]:
    width = max((len(row) for row in rows), default=0)
    if width > _MAX_TABLE_COLUMNS:
        raise ValueError(f"{label}列数 {width} 超过限制 {_MAX_TABLE_COLUMNS}，未执行审查")
    if width <= 0:
        return [], len(rows) > 0, 0
    header_limit = max(header_rows)
    row_limit = min(_MAX_TABLE_ROWS + header_limit, header_limit + max(1, cell_budget // width))
    bounded = rows[:row_limit]
    return bounded, len(rows) > row_limit, len(bounded) * width


def _load_frame(file_index: int, sheet_name: str = "", header_row: int = 1,
                header_rows: list[int] | None = None) -> tuple[Any, dict[str, Any]]:
    item = _attachment(file_index)
    name = str(item.get("name") or "table")
    suffix = Path(name).suffix.lower()
    data = item.get("data") or b""
    if suffix in {".csv", ".tsv"}:
        frame, encoding, raw_preview, used_header_rows, truncated = _decode_csv(data, header_row, header_rows)
        meta = {"name": name, "sheet": "", "encoding": encoding, "delimiter": "auto",
                "raw_preview": raw_preview, "header_row": max(used_header_rows),
                "header_rows": used_header_rows, "truncated": truncated}
    elif suffix in {".xlsx", ".xlsm", ".xls", ".xlsb"}:
        try:
            used_header_rows = sorted(set(header_rows or [header_row]))
            if not used_header_rows or any(row < 1 or row > 20 for row in used_header_rows):
                raise ValueError("header_row 必须在 1 到 20 之间")
            if suffix in {".xls", ".xlsb"}:
                try:
                    import python_calamine
                    reader = python_calamine.CalamineWorkbook.from_filelike(io.BytesIO(data))
                    names = reader.sheet_names
                    chosen = sheet_name.strip() if sheet_name.strip() in names else (names[0] if names else "")
                    if not chosen:
                        raise ValueError("工作簿中没有可读取的工作表")
                    sheet = reader.get_sheet_by_name(chosen)
                    width = getattr(sheet, "width", None)
                    height = getattr(sheet, "height", None)
                    width = width() if callable(width) else width
                    height = height() if callable(height) else height
                    if not isinstance(width, int) or not isinstance(height, int):
                        raise ValueError("无法在读取前确认工作表尺寸；为避免无界内存占用，已停止解析")
                    if width > _MAX_TABLE_COLUMNS:
                        raise ValueError(f"列数 {width} 超过限制 {_MAX_TABLE_COLUMNS}")
                    row_limit = max(1, min(_MAX_TABLE_ROWS + max(used_header_rows) + 1,
                                           _MAX_TABLE_CELLS // max(1, width) + max(used_header_rows)))
                    if height > row_limit:
                        raise ValueError(
                            f"工作表有 {height:,} 行，当前 XLS/XLSB 安全解析上限为 {row_limit:,} 行；"
                            "请先拆分工作表后重试，未对该表执行审查")
                    data_rows = sheet.to_python()
                    if len(data_rows) > row_limit:
                        raise ValueError("工作表实际行数超出安全解析范围，未执行审查")
                except ImportError as exc:
                    raise ValueError(f"读取 {suffix.upper()} 需要安装 python-calamine；文件未被判定为无异常") from exc
            else:
                import openpyxl
                from zipfile import ZipFile
                from src.core.file_safety import validate_zip_bytes
                validate_zip_bytes(data, label=name)
                with ZipFile(io.BytesIO(data)) as archive:
                    if "xl/workbook.xml" not in archive.namelist():
                        raise ValueError("工作簿缺少必要的 OOXML 结构")
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True,
                                            keep_vba=False)
                names = wb.sheetnames
                chosen = sheet_name.strip() if sheet_name.strip() in names else next(
                    (s for s in names if wb[s].max_row and wb[s].max_column), names[0] if names else "")
                if not chosen:
                    wb.close()
                    raise ValueError("工作簿中没有可读取的工作表")
                worksheet = wb[chosen]
                width = int(worksheet.max_column or 0)
                height = int(worksheet.max_row or 0)
                if width > _MAX_TABLE_COLUMNS:
                    wb.close()
                    raise ValueError(f"列数 {width} 超过限制 {_MAX_TABLE_COLUMNS}")
                row_limit = max(1, min(_MAX_TABLE_ROWS + max(used_header_rows) + 1,
                                       _MAX_TABLE_CELLS // max(1, width) + max(used_header_rows)))
                data_rows = list(worksheet.iter_rows(
                    min_row=1, max_row=min(height, row_limit), max_col=max(1, width),
                    values_only=True)) if height and width else []
                wb.close()
            data_row_limit = min(_MAX_TABLE_ROWS, _MAX_TABLE_CELLS // max(1, width))
            if len(data_rows) < max(used_header_rows):
                frame = _pd().DataFrame()
            else:
                width = max((len(row) for row in data_rows), default=0)
                if width > _MAX_TABLE_COLUMNS:
                    raise ValueError(f"列数 {width} 超过限制 {_MAX_TABLE_COLUMNS}")
                header = _normalize_headers(data_rows, used_header_rows, width)
                data_row_limit = min(_MAX_TABLE_ROWS, _MAX_TABLE_CELLS // max(1, len(header)))
                frame = _pd().DataFrame(
                    data_rows[max(used_header_rows):max(used_header_rows) + data_row_limit],
                    columns=header)
            meta = {"name": name, "sheet": chosen, "sheets": names, "encoding": suffix[1:],
                    "header_row": max(used_header_rows), "header_rows": used_header_rows,
                    "raw_preview": [[str(v)[:120] if v is not None else "" for v in row]
                                    for row in data_rows[:8]]}
            meta["truncated"] = (height > row_limit) or (
                len(data_rows) > max(used_header_rows) + data_row_limit)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"无法读取工作簿：{exc}") from exc
    elif suffix == ".pdf":
        try:
            import pdfplumber
            all_tables = []
            cell_budget = _MAX_TABLE_CELLS
            extraction_truncated = False
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = pdf.pages[:200]
                extraction_truncated = len(pdf.pages) > len(pages)
                for page_no, page in enumerate(pages, start=1):
                    for table_no, rows in enumerate(page.extract_tables() or [], start=1):
                        if rows and any(any(str(cell or "").strip() for cell in row) for row in rows):
                            if len(all_tables) >= _MAX_EXTRACTED_TABLES or cell_budget <= 0:
                                extraction_truncated = True
                                break
                            label = f"{name}!p{page_no}_t{table_no}"
                            bounded_rows, clipped, used_cells = _bound_raw_table(
                                rows, sorted(set(header_rows or [header_row])), cell_budget, label)
                            all_tables.append((label, bounded_rows))
                            cell_budget -= used_cells
                            extraction_truncated = extraction_truncated or clipped
                    if extraction_truncated and (len(all_tables) >= _MAX_EXTRACTED_TABLES or cell_budget <= 0):
                        break
            matching = next(((label, rows) for label, rows in all_tables
                             if label == sheet_name.strip()), None)
            if not all_tables:
                raise ValueError("PDF 中没有可提取的结构化表格；扫描件/OCR 未覆盖，审查未完成")
            chosen, rows = matching or all_tables[0]
            used_header_rows = sorted(set(header_rows or [header_row]))
            if any(row < 1 or row > 20 for row in used_header_rows):
                raise ValueError("header_row 必须在 1 到 20 之间")
            width = max((len(row) for row in rows), default=0)
            header = _normalize_headers(rows, used_header_rows, width)
            data_row_limit = min(_MAX_TABLE_ROWS, _MAX_TABLE_CELLS // max(1, width))
            frame = _pd().DataFrame(
                [(row + [None] * width)[:width]
                 for row in rows[max(used_header_rows):max(used_header_rows) + data_row_limit]],
                columns=header)
            meta = {"name": name, "sheet": chosen, "sheets": [label for label, _ in all_tables],
                    "encoding": "pdf-table", "header_row": max(used_header_rows),
                    "header_rows": used_header_rows, "raw_preview": rows[:8],
                    "truncated": extraction_truncated or
                                 len(rows) > max(used_header_rows) + data_row_limit}
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"无法提取 PDF 表格：{exc}") from exc
    elif suffix == ".docx":
        try:
            from src.core.file_safety import validate_zip_bytes
            validate_zip_bytes(data, label=name)
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                document_info = archive.getinfo("word/document.xml")
                if document_info.file_size > 16 * 1024 * 1024:
                    raise ValueError("Word 正文 XML 超过 16 MB，已停止解析")
                root = ET.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            word_tables = []
            cell_budget = _MAX_TABLE_CELLS
            extraction_truncated = False
            for number, table in enumerate(root.findall(".//w:tbl", ns), start=1):
                rows = []
                for row in table.findall("./w:tr", ns):
                    cells = row.findall("./w:tc", ns)
                    if len(cells) > _MAX_TABLE_COLUMNS:
                        raise ValueError(f"Word 表格 {number} 列数超过限制 {_MAX_TABLE_COLUMNS}，未执行审查")
                    rows.append(["".join(t.text or "" for t in cell.findall(".//w:t", ns))
                                 for cell in cells])
                    if sum(len(item) for item in rows) > _MAX_TABLE_CELLS:
                        extraction_truncated = True
                        break
                if rows:
                    if len(word_tables) >= _MAX_EXTRACTED_TABLES or cell_budget <= 0:
                        extraction_truncated = True
                        break
                    label = f"{name}!t{number}"
                    bounded_rows, clipped, used_cells = _bound_raw_table(
                        rows, sorted(set(header_rows or [header_row])), cell_budget, label)
                    word_tables.append((label, bounded_rows))
                    cell_budget -= used_cells
                    extraction_truncated = extraction_truncated or clipped
                    if cell_budget <= 0:
                        extraction_truncated = True
                        break
            if not word_tables:
                raise ValueError("Word 中没有可提取的结构化表格，审查未完成")
            matching = next(((label, rows) for label, rows in word_tables
                             if label == sheet_name.strip()), None)
            chosen, rows = matching or word_tables[0]
            used_header_rows = sorted(set(header_rows or [header_row]))
            width = max((len(row) for row in rows), default=0)
            header = _normalize_headers(rows, used_header_rows, width)
            data_row_limit = min(_MAX_TABLE_ROWS, _MAX_TABLE_CELLS // max(1, width))
            frame = _pd().DataFrame(
                [(row + [None] * width)[:width]
                 for row in rows[max(used_header_rows):max(used_header_rows) + data_row_limit]],
                columns=header)
            meta = {"name": name, "sheet": chosen, "sheets": [label for label, _ in word_tables],
                    "encoding": "docx-table", "header_row": max(used_header_rows),
                    "header_rows": used_header_rows, "raw_preview": rows[:8],
                    "truncated": extraction_truncated or
                                 len(rows) > max(used_header_rows) + data_row_limit}
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"无法提取 Word 表格：{exc}") from exc
    else:
        raise ValueError("当前审查支持 CSV/TSV/XLSX/XLS/XLSM/PDF/DOCX")
    if frame.shape[1] > _MAX_TABLE_COLUMNS:
        raise ValueError(f"列数 {frame.shape[1]} 超过限制 {_MAX_TABLE_COLUMNS}")
    if frame.shape[0] > _MAX_TABLE_ROWS:
        frame = frame.iloc[:_MAX_TABLE_ROWS].copy()
        meta["truncated"] = True
    else:
        meta["truncated"] = bool(meta.get("truncated", False))
    return frame, meta


def _json_arg(raw: Any, expected: type, default):
    if isinstance(raw, expected):
        return raw
    if raw is None or not str(raw or "").strip():
        return default
    if not isinstance(raw, str):
        raise ValueError(f"参数需要 JSON {expected.__name__} 格式")
    value = json.loads(raw)
    if not isinstance(value, expected):
        raise ValueError(f"参数需要 JSON {expected.__name__} 格式")
    return value


def _parse_header_rows(raw: str, default_header_row: int) -> list[int]:
    rows = _json_arg(raw, list, None) if raw else None
    if rows is None:
        rows = [default_header_row]
    if not rows or any(type(row) is not int for row in rows):
        raise ValueError("表头行参数必须是非空整数数组，例如 [1,2]")
    return sorted(set(rows))


def _profile(frame) -> dict[str, Any]:
    result = {"rows": int(frame.shape[0]), "columns": int(frame.shape[1]), "fields": []}
    if frame.empty or frame.shape[1] == 0:
        return result
    for col in frame.columns:
        series = frame[col]
        nonempty = series.dropna()
        field: dict[str, Any] = {
            "name": str(col), "inferred_type": str(series.dtype),
            "missing_count": int(series.isna().sum()),
            "missing_fraction": round(float(series.isna().mean()), 4),
            "unique_count": int(series.nunique(dropna=True)),
            "samples": [str(x)[:100] for x in nonempty.drop_duplicates().head(_MAX_SAMPLES).tolist()],
        }
        numeric = _pd().to_numeric(series, errors="coerce").dropna()
        source_nonempty = int(series.notna().sum())
        if len(numeric):
            field.update(numeric_count=int(len(numeric)), min=float(numeric.min()),
                         median=float(numeric.median()), max=float(numeric.max()),
                         numeric_fraction=round(len(numeric) / max(source_nonempty, 1), 4))
        result["fields"].append(field)
    result["sample_rows"] = frame.head(5).where(frame.notna(), None).to_dict(orient="records")
    return result


@tool
def inspect_uploaded_table(file_index: int = 0, sheet_name: str = "", header_row: int = 1,
                           header_rows_json: str = "") -> str:
    """查看本轮 CSV/TSV/Excel/PDF/Word 结构化表格的字段样例、基础统计及原始前几行。可用 header_row 选择表头，或 header_rows_json='[2,3]' 合并多层表头。先检查字段语义，再决定是否审查。"""
    header_rows = _parse_header_rows(header_rows_json, header_row)
    frame, meta = _load_frame(file_index, sheet_name, header_row, header_rows)
    profile = _profile(frame)
    profile.update(meta)
    profile["notice"] = "字段含义由模型结合列名与样例推断；表格内容仅限本轮，未写入知识库或长期记忆。"
    return json.dumps(profile, ensure_ascii=False, default=str)


def _audit_frame(frame, mapping: dict[str, str], groups: list[str], values: list[str],
                 row_offset: int = 2, integer_data_confirmed: bool = False,
                 sd_kind: str = ""):
    findings: list[dict[str, Any]] = []
    issue_rows: dict[int, set[str]] = {}
    def add(severity, check, field, detail, row=None):
        finding = {"severity": severity, "check": check, "field": str(field or ""),
                   "row": int(row) + row_offset if row is not None else "", "detail": str(detail)}
        findings.append(finding)
        if row is not None:
            issue_rows.setdefault(int(row), set()).add(check)

    if frame.empty:
        add("error", "empty_table", "", "没有数据行可供检查")
        return findings, issue_rows
    for col in frame.columns:
        series = frame[col]
        missing = int(series.isna().sum())
        if missing:
            add("info", "missing_values", col,
                f"{missing}/{len(series)} 行为空（{missing / len(series):.1%}）")
        if series.nunique(dropna=True) <= 1 and len(series.dropna()) > 0:
            add("warning", "constant_field", col, "非空值只有一个唯一取值；请确认该字段是否有信息量")
        numeric = _pd().to_numeric(series, errors="coerce")
        valid = numeric.dropna()
        nonempty_count = int(series.notna().sum())
        failed_numeric = nonempty_count - len(valid)
        if len(valid) and failed_numeric:
            add("review", "mixed_numeric_text", col,
                f"{failed_numeric} 个非空值无法解析为数值；可能混有单位、符号或文本类别")
        if len(valid) >= 8 and len(valid) / max(1, len(series.dropna())) >= 0.8:
            valid_groups = [c for c in groups if c in frame.columns]
            partitions = (frame.groupby(valid_groups, dropna=False).groups.values()
                          if valid_groups else [frame.index])
            for partition in partitions:
                local = numeric.loc[list(partition)].dropna()
                if len(local) < 8:
                    continue
                q1, q3 = float(local.quantile(.25)), float(local.quantile(.75))
                iqr = q3 - q1
                if iqr <= 0:
                    continue
                low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
                for idx in local[(local < low) | (local > high)].index[:10000]:
                    add("review", "iqr_outlier", col,
                        f"值 {numeric.loc[idx]:g} 超出探索性 IQR 范围 [{low:g}, {high:g}]；不代表错误，建议结合实验背景复核", idx)
    duplicate_mask = frame.duplicated(keep=False)
    for idx in frame.index[duplicate_mask][:10000]:
        add("review", "duplicate_row", "", "存在完全重复行，请确认是否为重复录入或合法重复测量", idx)
    # Numeric-forensics cues adapted from PaperConan's published detector
    # taxonomy. These are review signals, never judgments about intent.
    numeric_cols = []
    for col in frame.columns[:30]:
        numeric = _pd().to_numeric(frame[col], errors="coerce")
        if int(numeric.notna().sum()) >= 5 and numeric.notna().mean() >= .8:
            numeric_cols.append((col, numeric))
    for i, (left_name, left) in enumerate(numeric_cols):
        for right_name, right in numeric_cols[i + 1:]:
            valid = left.notna() & right.notna()
            if int(valid.sum()) < 5:
                continue
            positions = frame.index[valid]
            if len(positions) > 20_000:
                positions = positions[:20_000]
            a = left.loc[positions].astype(float); b = right.loc[positions].astype(float)
            pair_count = len(a)
            detail = ""
            if _pd().Series(a.to_numpy()).equals(_pd().Series(b.to_numpy())):
                detail = f"{pair_count} 个共同非空位置逐值完全相同"
                check = "identical_columns"
            else:
                delta = b - a
                scale = max(float(a.abs().max()), float(b.abs().max()), 1.0)
                tolerance = max(1e-9, scale * 1e-9)
                if float(delta.max() - delta.min()) <= tolerance:
                    detail = f"两列在 {pair_count} 个共同位置的差值近似固定（约 {float(delta.median()):g}）"
                    check = "constant_offset_columns"
                elif float(a.abs().max()) > tolerance:
                    ratio_mask = a.abs() > tolerance
                    ratios = b[ratio_mask] / a[ratio_mask]
                    if len(ratios) >= 5 and float(ratios.max() - ratios.min()) <= max(1e-8, float(ratios.abs().max()) * 1e-8):
                        detail = f"两列在 {len(ratios)} 个共同位置的比值近似固定（约 {float(ratios.median()):g}）"
                        check = "constant_ratio_columns"
                if not detail:
                    try:
                        slope, intercept = np.polyfit(a.to_numpy(), b.to_numpy(), 1)
                        residual = b.to_numpy() - (slope * a.to_numpy() + intercept)
                        if float(np.max(np.abs(residual))) <= tolerance:
                            detail = f"两列近似满足线性关系 y={slope:g}x+{intercept:g}（{pair_count} 点）"
                            check = "exact_linear_columns"
                    except Exception:
                        pass
            if detail:
                add("review", check, f"{left_name} ↔ {right_name}",
                    detail + "；可能来自换算、归一化、公式或共享对照，需结合方法学确认")
    # Repeated fine-grained decimal endings are weak signals and are only
    # surfaced when the column has enough distinct continuous-looking values.
    for col, numeric in numeric_cols:
        vals = numeric.dropna().astype(float)
        precise = vals[(vals - vals.round(2)).abs() > 1e-6]
        if len(precise) < 12 or precise.nunique() < 8:
            continue
        endings = precise.map(lambda x: f"{abs(x):.6f}".split(".")[1][-2:])
        counts = endings.value_counts(normalize=True)
        common = counts[counts >= .35]
        if not common.empty:
            add("review", "decimal_ending_concentration", col,
                "小数末两位存在集中（" + ", ".join(f".{k}: {v:.0%}" for k, v in common.head(4).items()) +
                "）；均值、比例、舍入、固定分母及仪器精度也会产生此现象")
    if groups:
        absent = [c for c in groups if c not in frame.columns]
        if absent:
            add("warning", "invalid_group_mapping", ",".join(absent), "模型提供的分组字段不在表格中；未执行分组平衡检查")
        else:
            counts = frame.groupby(groups, dropna=False).size()
            if len(counts) > 1 and counts.nunique() > 1:
                distribution = "; ".join(f"{k}: n={v}" for k, v in list(counts.items())[:80])
                add("review", "unequal_group_size", ",".join(groups),
                    "分组样本数不一致（这可能是有意设计）；" + distribution)
    valid_values = [c for c in values if c in frame.columns]
    if values and len(valid_values) != len(values):
        add("warning", "invalid_value_mapping", ",".join(set(values) - set(valid_values)),
            "指定的测量字段未全部匹配列名")
    # column_mapping_json maps semantic role -> literal source column name.
    role_cols = {str(role): str(column) for role, column in mapping.items()
                 if column in frame.columns}
    if "replicate" in role_cols and groups:
        rep = role_cols["replicate"]
        keys = [c for c in groups if c in frame.columns]
        rep_counts = frame.groupby(keys, dropna=False)[rep].nunique(dropna=True)
        if len(rep_counts) > 1 and rep_counts.nunique() > 1:
            add("review", "unequal_replicates", rep,
                "不同实验组的重复编号数不一致；尚不能判断为设计问题，请核对生物学/技术重复定义")
    if integer_data_confirmed and {"mean", "n"}.issubset(role_cols):
        means = _pd().to_numeric(frame[role_cols["mean"]], errors="coerce")
        ns = _pd().to_numeric(frame[role_cols["n"]], errors="coerce")
        if "sd" in role_cols and sd_kind not in {"sample", "population"}:
            add("info", "grim_sd_skipped", role_cols["sd"],
                "SD一致性检查未执行：请先确认报告的是样本SD还是总体SD；均值与样本量检查仍可执行")
        for idx in frame.index[means.notna() & ns.notna()]:
            mean = float(means.loc[idx]); n = float(ns.loc[idx])
            if n <= 0 or abs(n - round(n)) > 1e-9:
                continue
            n = int(round(n))
            if n > 1 and abs(mean * n - round(mean * n)) > .02:
                add("review", "integer_mean_inconsistent", role_cols["mean"],
                    f"若为 {n} 个整数观测的算术均值，均值×n={mean*n:g} 不接近整数；请核对数据粒度、舍入和 n 定义", idx)
            if "sd" in role_cols and sd_kind in {"sample", "population"} and n > 1:
                sd_value = _pd().to_numeric(frame[role_cols["sd"]], errors="coerce").get(idx)
                if sd_value is None or _pd().isna(sd_value) or float(sd_value) < 0:
                    continue
                factor = (n - 1) if sd_kind == "sample" else n
                squares = factor * float(sd_value) ** 2 + n * mean ** 2
                if abs(squares - round(squares)) > .08:
                    add("review", "integer_summary_inconsistent", role_cols["sd"],
                        f"若为整数观测且按{('样本' if sd_kind == 'sample' else '总体')}SD计算，推得平方和 {squares:g} 不接近整数；请核对舍入与统计定义", idx)
    if not groups:
        add("info", "design_checks_skipped", "", "未确认实验分组字段，只执行通用数据检查")
    return findings[:20000], issue_rows


def _write_audit_xlsx(filename: str, meta: dict, findings: list[dict]):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "审查摘要"
    ws.append(["项目", "结果"])
    ws.append(["原文件", _safe_spreadsheet_value(meta.get("name", ""))])
    ws.append(["工作表", _safe_spreadsheet_value(meta.get("sheet", ""))])
    ws.append(["行数", meta.get("rows", 0)])
    ws.append(["列数", meta.get("columns", 0)])
    ws.append(["异常与提示数", len(findings)])
    ws.append(["说明", "离群与组间不平衡为复核提示，不等于错误；原始数据未修改。"])
    detail = wb.create_sheet("发现明细")
    headers = ["级别", "检查项", "文件", "工作表", "位置", "字段", "说明", "证据与复核上下文"]
    detail.append(headers)
    for item in findings:
        evidence = item.get("evidence", "")
        if not isinstance(evidence, str):
            evidence = json.dumps(evidence, ensure_ascii=False, default=str)
        detail.append([_safe_spreadsheet_value(item.get("severity")),
                       _safe_spreadsheet_value(item.get("check")),
                       _safe_spreadsheet_value(item.get("file")),
                       _safe_spreadsheet_value(item.get("sheet")),
                       _safe_spreadsheet_value(item.get("location") or item.get("row")),
                       _safe_spreadsheet_value(item.get("field")),
                       _safe_spreadsheet_value(item.get("detail")),
                       _safe_spreadsheet_value(evidence)])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(70, max(12, max(len(str(cell.value or "")) for cell in column[:200]) + 2))
            sheet.column_dimensions[letter].width = width
    if meta.get("scan_errors"):
        errors = wb.create_sheet("未完成项")
        errors.append(["文件或对象", "状态"])
        for value in meta["scan_errors"]:
            errors.append([_safe_spreadsheet_value(value), "未完成；不能视作通过"])
        errors.freeze_panes = "A2"
        errors.column_dimensions["A"].width = 90
        errors.column_dimensions["B"].width = 28
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _extract_attachment_text(item: dict[str, Any]) -> str:
    name = str(item.get("name") or "")
    data = item.get("data") or b""
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            text = "\n".join(doc.load_page(i).get_text("text") for i in range(min(len(doc), 200)))
            doc.close()
            return text
        if suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                root = ET.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            return "\n".join("".join(node.text or "" for node in p.findall(".//w:t", ns))
                             for p in root.findall(".//w:p", ns))
        if suffix in {".csv", ".tsv"}:
            return data.decode("utf-8-sig", errors="replace")[:1_000_000]
    except Exception as exc:
        return f"[文本提取失败：{type(exc).__name__}]"
    return ""


def _normalize_doi(value: str) -> str:
    doi = str(value or "").strip().lower()
    doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", doi, flags=re.I)
    doi = re.sub(r"^doi\s*:\s*", "", doi, flags=re.I)
    return doi.rstrip(".,;:，。；）)]}")


def _title_from_metadata(record: dict[str, Any], provider: str) -> str:
    if provider == "crossref":
        values = record.get("title") or []
        return str(values[0]) if values else ""
    return str(record.get("titles", [{}])[0].get("title") or "")


def _metadata_candidate(doi: str, timeout: float = 5.0) -> tuple[str, dict | None, str]:
    """Look up DOI metadata in Crossref and DataCite; return status and record."""
    import httpx
    saw_network_failure = False
    for provider, url in (
        ("crossref", f"https://api.crossref.org/works/{quote(doi, safe='/')}") ,
        ("datacite", f"https://api.datacite.org/dois/{quote(doi, safe='/')}") ,
    ):
        try:
            response = httpx.get(url, timeout=timeout, follow_redirects=True,
                                 headers={"User-Agent": "CitrusQA research audit/1.0"})
            if response.status_code == 200:
                payload = response.json()
                record = payload.get("message") if provider == "crossref" else payload.get("data", {}).get("attributes")
                if isinstance(record, dict):
                    return "found", record, provider
            elif response.status_code != 404:
                saw_network_failure = True
        except Exception:
            saw_network_failure = True
    return ("unverified" if saw_network_failure else "not_found_in_checked_registries"), None, "Crossref/DataCite"


def _metadata_by_title(title: str, timeout: float = 5.0) -> tuple[dict | None, float, str, bool]:
    import httpx
    try:
        response = httpx.get("https://api.crossref.org/works", params={
            "query.bibliographic": title[:300], "rows": 5,
        }, timeout=timeout, follow_redirects=True,
           headers={"User-Agent": "CitrusQA research audit/1.0"})
        response.raise_for_status()
        items = response.json().get("message", {}).get("items", [])
        candidates = []
        for record in items:
            candidate_title = _title_from_metadata(record, "crossref")
            score = SequenceMatcher(None, title.casefold(), candidate_title.casefold()).ratio()
            candidates.append((score, record))
        if not candidates:
            return None, 0.0, "Crossref", False
        score, record = max(candidates, key=lambda pair: pair[0])
        return record, score, "Crossref", False
    except Exception:
        return None, 0.0, "Crossref", True


def _verify_reference_record(raw: dict[str, Any]) -> dict[str, Any]:
    doi = _normalize_doi(raw.get("doi", ""))
    if not doi:
        title = str(raw.get("title") or "").strip()
        if not title:
            return {"doi": "", "status": "not_checked", "detail": "缺少 DOI 与标题；需要人工核对"}
        record, score, provider, network_error = _metadata_by_title(title)
        if not record or score < .82:
            return {"doi": "", "status": "unverified" if network_error else "not_found_or_low_match",
                    "input_title": title,
                    "detail": "Crossref 查询失败，未完成核验。" if network_error else
                              "Crossref 未返回足够匹配的标题候选；不能据此断言文献不存在。"}
        issued = (record.get("issued", {}).get("date-parts") or [[""]])[0]
        year = str(issued[0]) if issued else ""
        authors = "; ".join(" ".join(filter(None, [a.get("given"), a.get("family")]))
                             for a in (record.get("author") or [])[:8])
        author_source = str(raw.get("authors") or "").casefold()
        author_mismatch = bool(author_source and authors and
                               not any(token.strip() and token.strip() in authors.casefold()
                                       for token in re.split(r"[,;，；]", author_source)[:4]))
        return {"doi": str(record.get("DOI") or ""), "provider": provider,
                "status": "title_candidate", "input_title": title,
                "metadata_title": _title_from_metadata(record, "crossref"),
                "input_authors": raw.get("authors", ""), "metadata_authors": authors,
                "input_year": str(raw.get("year") or "")[:4], "metadata_year": year,
                "metadata_journal": "; ".join(record.get("container-title") or []),
                "title_similarity": round(score, 3),
                "detail": "标题匹配候选，不是 DOI 精确核验；请人工确认作者、期刊与年份。"
                          + ("作者字段可能不一致。" if author_mismatch else "")}
    status, record, provider = _metadata_candidate(doi)
    row = {"doi": doi, "provider": provider, "status": status,
           "input_title": str(raw.get("title") or ""), "input_authors": str(raw.get("authors") or ""),
           "input_year": str(raw.get("year") or ""), "citation": str(raw.get("citation") or "")}
    if not record:
        row["detail"] = ("在本次查询的登记库中未找到记录；不能据此断言文献不存在。"
                         if status == "not_found_in_checked_registries" else
                         "联网失败或登记库暂不可用，未完成核验。")
        return row
    title = _title_from_metadata(record, provider)
    if provider == "crossref":
        authors = "; ".join(" ".join(filter(None, [a.get("given"), a.get("family")]))
                             for a in (record.get("author") or [])[:8])
        issued = (record.get("issued", {}).get("date-parts") or [[""]])[0]
        year = str(issued[0]) if issued else ""
        journal = "; ".join(record.get("container-title") or [])
        metadata_doi = str(record.get("DOI") or "")
    else:
        creators = record.get("creators") or []
        authors = "; ".join(str(a.get("name") or "") for a in creators[:8])
        year = str(record.get("publicationYear") or "")
        journal = str((record.get("container") or {}).get("title") or "")
        metadata_doi = str(record.get("doi") or doi)
    row.update(metadata_title=title, metadata_authors=authors, metadata_year=year,
               metadata_journal=journal, doi_matches=(_normalize_doi(metadata_doi) == doi))
    expected = str(raw.get("title") or "").strip()
    row["title_similarity"] = round(SequenceMatcher(None, expected.casefold(), title.casefold()).ratio(), 3) if expected else None
    author_tokens = [token.strip() for token in re.split(r"[,;，；]", row["input_authors"].casefold())[:4] if token.strip()]
    row["author_mismatch"] = bool(author_tokens and authors and
                                  not any(token in authors.casefold() for token in author_tokens))
    row["field_mismatch"] = bool(
        row["doi_matches"] is False or (expected and row["title_similarity"] < .82) or
        row["author_mismatch"] or (raw.get("year") and str(raw.get("year"))[:4] != year[:4]))
    row["status"] = "metadata_mismatch" if row["field_mismatch"] else "metadata_matched"
    return row


@tool
def verify_uploaded_references(file_index: int = 0, references_json: str = "") -> str:
    """核验本轮论文/参考文献附件中的 DOI 与登记元数据。可传 references_json=[{"doi":"...","title":"...","authors":"...","year":"..."}]，支持没有 DOI 的题名记录；不传时仅从附件自动提取 DOI。未检索到不等于文献不存在。"""
    supplied = _json_arg(references_json, list, [])
    files = get_request_tables()
    item = files[file_index] if 0 <= file_index < len(files) else {"name": "references", "data": b""}
    references = supplied
    if not references:
        if not item.get("data"):
            return "请提供带 DOI 的参考文献清单，或上传论文/文献表；当前请求没有可提取的来源。"
        text = _extract_attachment_text(item)
        if not text.strip():
            return "附件没有可搜索文本（可能是扫描版 PDF）；参考文献核验未完成。请上传可复制文本的 PDF/Word，或提供 DOI/题名。"
        seen = set()
        references = []
        for match in _DOI_PATTERN.findall(text):
            doi = _normalize_doi(match)
            if doi and doi not in seen:
                seen.add(doi)
                line = next((line.strip() for line in text.splitlines() if doi.lower() in line.lower()), "")
                references.append({"doi": doi, "citation": line[:500]})
    if not references:
        if text.startswith("[文本提取失败："):
            return f"{text}；附件核验未完成。请重新上传可复制文本的 PDF/Word，或提供参考文献 DOI/题名。"
        return ("附件文本已读取，但未自动提取到 DOI。请从正文参考文献中提取 DOI；无 DOI 条目请传入 references_json，"
                "格式为 [{\"title\":\"题名\",\"authors\":\"作者\",\"year\":\"年份\"}] 再核验。"
                "这不代表文献不存在。")
    from src.core.executors import AUDIT_NETWORK_EXECUTOR
    selected = [raw for raw in references[:100] if isinstance(raw, dict)]
    rows = []
    # Bound queued work per request as well as global active threads.
    for start in range(0, len(selected), 16):
        rows.extend(AUDIT_NETWORK_EXECUTOR.map(
            _verify_reference_record, selected[start:start + 16]))
    if len(references) > len(selected):
        rows.append({"doi": "", "status": "not_checked", "detail": f"另有 {len(references) - len(selected)} 条未核验（单次上限 100 条）。"})
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "DOI核验"
    headers = ["DOI", "状态", "登记库", "输入标题", "登记标题", "输入作者", "登记作者", "输入年份", "登记年份", "期刊", "标题相似度", "说明"]
    ws.append(headers)
    for row in rows:
        ws.append([_safe_spreadsheet_value(row.get(k, "")) for k in
                   ("doi", "status", "provider", "input_title", "metadata_title", "input_authors",
                    "metadata_authors", "input_year", "metadata_year", "metadata_journal", "title_similarity", "detail")])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO(); wb.save(out)
    filename = _safe_name(Path(item.get("name", "paper")).stem, "文献") + "_DOI核验.xlsx"
    url = _save_artifact(filename, out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    counts = {}
    for row in rows: counts[row["status"]] = counts.get(row["status"], 0) + 1
    return (f"已核验 {len(rows)} 条 DOI：" + "、".join(f"{key} {value} 条" for key, value in counts.items()) +
            "。登记元数据不证明文献内容支持某项结论；未查到或网络失败不会判作虚假文献。\n\n"
            f"[下载 DOI 核验表]({url})")


def _audit_category_for_kind(kind: str) -> str:
    known = {check: label for label, checks in _AUDIT_CATEGORIES for check in checks}
    if kind in known:
        return known[kind]
    lowered = kind.lower()
    if any(token in lowered for token in ("duplicate", "reuse", "identical", "repeat", "block_dup")):
        return "重复值与数据复用"
    if any(token in lowered for token in ("outlier", "anomal", "range")):
        return "异常值"
    if any(token in lowered for token in ("grim", "mean", "sd", "summary")):
        return "均值与标准差自洽性"
    if any(token in lowered for token in ("digit", "decimal", "tail", "ending")):
        return "数字分布与小数尾"
    if any(token in lowered for token in ("row", "relation", "offset", "ratio", "linear", "sum", "diff", "progression")):
        return "行列关系与局部数值模式"
    if any(token in lowered for token in ("missing", "empty", "mixed", "invalid", "constant_field")):
        return "缺失值与数据格式"
    if any(token in lowered for token in ("group", "replicate", "design")):
        return "组别与重复数"
    return "其他审查线索"


def _audit_report_subject(source_name: str) -> str:
    stem = Path(source_name).stem
    stem = re.sub(r"^(?:[0-9a-f]{20,}[_-]+)+", "", stem, flags=re.IGNORECASE)
    stem = re.sub(
        r"(?:[_\-\s]*(?:科研数据审查|数据审查|审查结果|仪器导出|仪器数据|数据导出|科研数据|审查|导出|报告))+$",
        "", stem,
        flags=re.IGNORECASE,
    )
    stem = re.sub(r"^[_\-\s]+|[_\-\s]+$", "", stem)
    return _safe_name(stem[:36], "科研数据")


def _render_audit_html(findings: list[dict[str, Any]], *, title: str, scan: dict[str, Any],
                       inspected_rows: int, inspected_cols: int,
                       skipped: list[str]) -> bytes:
    """Render one concise, self-contained HTML report with escaped user data."""
    grouped: dict[str, list[dict[str, Any]]] = {label: [] for label, _ in _AUDIT_CATEGORIES}
    not_run: dict[str, list[str]] = {
        "组别与重复数": [],
        "均值与标准差自洽性": [],
    }
    for finding in findings:
        kind = str(finding.get("check") or finding.get("kind") or "")
        if kind == "design_checks_skipped":
            not_run["组别与重复数"].append(str(finding.get("detail") or kind))
            continue
        if kind == "grim_sd_skipped":
            not_run["均值与标准差自洽性"].append(str(finding.get("detail") or kind))
            continue
        label = _audit_category_for_kind(kind)
        grouped[label].append(finding)

    coverage = scan.get("coverage") or {}
    sheets = len((scan.get("scan_stats") or {}).get("sheets") or [])
    n_files = int(scan.get("n_files") or coverage.get("files_succeeded") or 0)
    status = str(scan.get("scan_status") or "partial")
    total_findings = sum(map(len, grouped.values()))
    esc = lambda value: html.escape(str(value or ""), quote=True)
    status_text = "已完成" if status == "complete" else "部分完成"
    rows: list[str] = []
    coverage_notice = ""
    if status != "complete":
        limits = (coverage.get("limitations") or [])
        coverage_notes = [str(item) for item in [*limits, *skipped] if item]
        note = "未检查范围不能写作“无异常”。"
        if coverage_notes:
            note += " 未完成项：" + "；".join(dict.fromkeys(coverage_notes[:3]))
        coverage_notice = f'<div class="notice"><strong>覆盖不完整</strong><span>{esc(note)}</span></div>'

    for label, _ in _AUDIT_CATEGORIES:
        items = grouped[label]
        if not items and not_run.get(label):
            result = "未执行：" + "；".join(list(dict.fromkeys(not_run[label]))[:2])
            rows.append(f'<tr><th>{esc(label)}</th><td><span class="state skipped">未执行</span> {esc(result)}</td></tr>')
            continue
        if not items and not not_run.get(label):
            rows.append(f'<tr><th>{esc(label)}</th><td><span class="state clear">无</span><span class="muted">已检查范围内</span></td></tr>')
            continue

        # Fold repeated hits by rule and table location. Show a few concrete
        # examples so large scans remain readable while retaining traceability.
        buckets: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
        for finding in items:
            key = (
                str(finding.get("check") or finding.get("kind") or "复核线索"),
                str(finding.get("file") or ""),
                str(finding.get("sheet") or ""),
            )
            buckets.setdefault(key, []).append(finding)
        examples: list[str] = []
        bucket_items = list(buckets.items())
        for (kind, file_name, sheet), hits in bucket_items[:3]:
            label_parts = [part for part in (kind or "复核线索", file_name, sheet) if part]
            fields = list(dict.fromkeys(str(hit.get("field") or "") for hit in hits if hit.get("field")))[:4]
            locations = list(dict.fromkeys(
                str(hit.get("location") or hit.get("range") or "") for hit in hits
                if hit.get("location") or hit.get("range")
            ))[:3]
            details = list(dict.fromkeys(
                str(hit.get("detail") or hit.get("rule") or "") for hit in hits
                if hit.get("detail") or hit.get("rule")
            ))[:2]
            example = " · ".join(label_parts) + f"：{len(hits)} 条"
            if fields:
                example += "；字段 " + "、".join(fields)
            if locations:
                example += "；位置 " + "、".join(locations)
            if details:
                detail_text = "；".join(details)
                if len(detail_text) > 140:
                    detail_text = detail_text[:137].rstrip() + "…"
                example += "；" + detail_text
            examples.append(example.replace("|", "／").replace("\n", " "))
        extra_groups = max(0, len(bucket_items) - len(examples))
        result_parts = [f'<strong>{len(items)} 条待复核</strong>']
        if examples:
            result_parts.append("<ul>" + "".join(f"<li>{esc(example)}</li>" for example in examples) + "</ul>")
        if extra_groups:
            result_parts.append(f'<span class="muted">另有 {extra_groups} 组同类线索，已计入数量</span>')
        if not_run.get(label):
            result_parts.append(f'<p><span class="state skipped">未执行</span> {esc("；".join(list(dict.fromkeys(not_run[label]))[:2]))}</p>')
        rows.append(f'<tr><th>{esc(label)}</th><td>{"".join(result_parts)}</td></tr>')

    source_label = esc(title)
    html_report = f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{source_label}报告</title><style>
:root{{color-scheme:light;--ink:#202a24;--muted:#6c756f;--line:#e3e9e4;--green:#2f7651;--pale:#f4f8f4;--amber:#9a5b08}}
*{{box-sizing:border-box}}body{{margin:0;background:#f5f7f5;color:var(--ink);font:15px/1.65 system-ui,-apple-system,"Segoe UI","Microsoft YaHei",sans-serif}}
main{{max-width:1020px;margin:36px auto;padding:0 20px}}.card{{background:#fff;border:1px solid var(--line);border-radius:16px;box-shadow:0 8px 28px #18351c0a;overflow:hidden}}
header{{padding:26px 30px 20px;border-bottom:1px solid var(--line)}}h1{{margin:0 0 14px;font-size:24px;letter-spacing:.02em}}.stats{{display:flex;flex-wrap:wrap;gap:9px}}
.stat{{padding:6px 11px;background:var(--pale);border:1px solid var(--line);border-radius:999px;font-size:13px}}.stat b{{color:var(--green)}}
.intro{{margin:14px 0 0;color:var(--muted);font-size:13px}}.notice{{margin:18px 30px 0;padding:12px 14px;border-radius:10px;background:#fff7e8;color:#754d15;border:1px solid #f1dfbf;display:grid;gap:3px}}
.table-wrap{{padding:18px 30px 26px;overflow-x:auto}}table{{width:100%;border-collapse:collapse}}th,td{{padding:14px 12px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}}
th{{width:190px;font-weight:650;white-space:nowrap}}tr:last-child th,tr:last-child td{{border-bottom:0}}td strong{{display:inline-block;margin:2px 0 4px}}
ul{{margin:5px 0 0;padding-left:20px}}li+li{{margin-top:5px}}.muted{{color:var(--muted);font-size:13px}}.state{{display:inline-block;padding:1px 8px;border-radius:999px;font-size:12px;font-weight:650;margin-right:7px}}
.clear{{color:var(--green);background:#edf7f0}}.skipped{{color:var(--amber);background:#fff2d8}}footer{{padding:14px 30px;border-top:1px solid var(--line);color:var(--muted);font-size:12px}}
@media(max-width:640px){{main{{margin:14px auto;padding:0 10px}}header,.table-wrap{{padding-left:18px;padding-right:18px}}.notice{{margin-left:18px;margin-right:18px}}th{{width:130px;white-space:normal}}}}
</style></head><body><main><article class="card"><header><h1>{source_label}报告</h1>
<div class="stats"><span class="stat">状态：<b>{esc(status_text)}</b></span><span class="stat">文件/工作表：<b>{n_files}/{sheets}</b></span>
<span class="stat">检查范围：<b>{inspected_rows} 行 × {inspected_cols} 列</b></span><span class="stat">待复核：<b>{total_findings} 条</b></span></div>
<p class="intro">规则命中仅为复核线索，不代表数据错误或研究不端。</p></header>
{coverage_notice}
<div class="table-wrap"><table><thead><tr><th>检查类别</th><th>结果与位置</th></tr></thead><tbody>{"".join(row for row in rows if row.startswith("<tr>") )}</tbody></table></div>
<footer>本报告由数据规则扫描生成。请结合原始记录、实验设计与文稿说明人工复核。</footer></article></main></body></html>'''
    return html_report.encode("utf-8")


@tool
def audit_uploaded_data(file_indexes_json: str | list[int] = "",
                        column_mapping_json: str | dict[str, str] = "{}",
                        group_columns_json: str | list[str] = "[]",
                        value_columns_json: str | list[str] = "[]",
                        header_row: int = 1, header_rows_json: str | list[int] = "",
                        integer_data_confirmed: bool = False, sd_kind: str = "",
                        include_report: bool = True) -> str:
    """统一审查本轮上传的一份或多份表格/论文附件。先 inspect_uploaded_table 理解复杂表头和字段；调用 PaperConan 数值检测并补充数据质量检查。支持跨工作表/文件、行列/局部/块级/小数尾/末位模式及整数摘要统计。所有命中只是复核线索，不推断意图、不修改原件；每次只生成一份按类别汇总的 HTML 异常报告，并标明完整/部分覆盖状态。"""
    audit_started = time.perf_counter()
    tables = get_request_tables()
    requested = _json_arg(file_indexes_json, list, list(range(len(tables))))
    if any(type(index) is not int for index in requested):
        raise ValueError("file_indexes_json 必须是整数数组")
    if not requested:
        return "审查未启动：当前轮没有选择附件。"
    mapping = _json_arg(column_mapping_json, dict, {})
    groups = _json_arg(group_columns_json, list, [])
    values = _json_arg(value_columns_json, list, [])
    if any(not isinstance(k, str) or not isinstance(v, str) for k, v in mapping.items()):
        raise ValueError("column_mapping_json 需要字符串到字符串的字段映射")
    if any(not isinstance(v, str) for v in groups + values):
        raise ValueError("group_columns_json 和 value_columns_json 需要字符串数组")
    header_rows = _parse_header_rows(header_rows_json, int(header_row))

    from src.research.paperconan_adapter import scan_request_tables
    scan_started = time.perf_counter()
    numeric = scan_request_tables(tables, requested)
    scan_ms = int((time.perf_counter() - scan_started) * 1000)
    scan = numeric["scan"]
    findings = list(numeric["findings"])
    skipped = list(scan.get("scan_errors") or [])
    inspected_rows = 0
    inspected_cols = 0
    quality_started = time.perf_counter()
    # Keep general table-quality checks that complement PaperConan. Relation and
    # digit-pattern duplicates are intentionally omitted here; PaperConan owns them.
    upstream_checks = {"identical_columns", "constant_offset_columns", "constant_ratio_columns",
                       "exact_linear_columns", "decimal_ending_concentration"}
    for index in requested:
        if index < 0 or index >= len(tables):
            raise ValueError(f"附件序号无效：{index}")
        item = tables[index]
        try:
            _, first_meta = _load_frame(index, header_row=int(header_row), header_rows=header_rows)
            sheet_names = first_meta.get("sheets") or [first_meta.get("sheet") or ""]
            for sheet in sheet_names:
                try:
                    frame, meta = _load_frame(index, sheet_name=str(sheet or ""),
                                               header_row=int(header_row), header_rows=header_rows)
                    inspected_rows += int(frame.shape[0])
                    inspected_cols += int(frame.shape[1])
                    base_findings, _ = _audit_frame(
                        frame, mapping, groups, values,
                        row_offset=max(meta.get("header_rows") or [1]) + 1,
                        integer_data_confirmed=bool(integer_data_confirmed),
                        sd_kind=str(sd_kind or "").lower())
                    source_name = str(item.get("name") or f"附件 {index + 1}")
                    source_sheet = str(meta.get("sheet") or "CSV")
                    for finding in base_findings:
                        if finding.get("check") in upstream_checks:
                            continue
                        finding.update(file=source_name, sheet=source_sheet,
                                       location=(f"第 {finding['row']} 行" if finding.get("row") else ""),
                                       evidence="", prefilter="")
                        findings.append(finding)
                except Exception as exc:
                    skipped.append(f"{item.get('name')} · {sheet}: {type(exc).__name__}: {exc}")
        except Exception as exc:
            skipped.append(f"{item.get('name')}: {type(exc).__name__}: {exc}")
    quality_ms = int((time.perf_counter() - quality_started) * 1000)
    status = str(scan.get("scan_status") or "partial")
    coverage = scan.get("coverage") or {}
    if skipped and status == "complete":
        status = "partial"
    source_names = [str(tables[index].get("name") or f"附件{index + 1}") for index in requested]
    display_name = "、".join(source_names[:3]) + (f" 等 {len(source_names)} 个文件" if len(source_names) > 3 else "")
    safe_skipped = [value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
                    for value in skipped]
    non_finding_statuses = {"design_checks_skipped", "grim_sd_skipped"}
    chat_findings = [finding for finding in findings
                     if str(finding.get("check") or finding.get("kind") or "") not in non_finding_statuses]

    # Always produce one user-facing HTML artifact. The structured scan remains
    # internal instead of becoming several separate downloads.
    report_subject = _audit_report_subject(source_names[0] if source_names else "")
    report_name = f"{report_subject}报告.html"
    scan["name"] = display_name
    render_started = time.perf_counter()
    html_report = _render_audit_html(findings, title=report_subject, scan=scan,
                                     inspected_rows=inspected_rows, inspected_cols=inspected_cols,
                                     skipped=safe_skipped)
    render_ms = int((time.perf_counter() - render_started) * 1000)
    artifact_started = time.perf_counter()
    report_url = _save_artifact(report_name, html_report, "text/html; charset=utf-8")
    artifact_ms = int((time.perf_counter() - artifact_started) * 1000)
    logger.info("[DataAudit] complete status=%s files=%d rows=%d columns=%d findings=%d "
                "scan_ms=%d quality_ms=%d render_ms=%d artifact_ms=%d total_ms=%d",
                status, len(requested), inspected_rows, inspected_cols, len(chat_findings),
                scan_ms, quality_ms, render_ms, artifact_ms,
                int((time.perf_counter() - audit_started) * 1000))
    report_links = f"\n\n[打开异常报告]({report_url})"
    if status == "complete":
        skipped_checks = len(findings) - len(chat_findings)
        if chat_findings:
            category_counts: dict[str, int] = {}
            for finding in chat_findings:
                kind = str(finding.get("check") or finding.get("kind") or "")
                category = _audit_category_for_kind(kind)
                category_counts[category] = category_counts.get(category, 0) + 1
            category_summary = "；".join(
                f"{label} {count} 条" for label, count in category_counts.items()
            )
            skipped_summary = f"；另有 {skipped_checks} 项检查未执行" if skipped_checks else ""
            summary = f"数据审查已完成：{category_summary}{skipped_summary}。详情见异常报告。"
        elif skipped_checks:
            summary = f"数据审查已完成，但有 {skipped_checks} 项检查未执行；这些项目不能视为无异常。详情见异常报告。"
        else:
            summary = "数据审查已完成，当前规则未发现待复核线索。详情见异常报告。"
    else:
        summary = "数据审查未完整完成，未完成范围不能判断为无异常。"
        if chat_findings:
            summary += f"已发现 {len(chat_findings)} 条待复核线索；详情见异常报告。"
        if safe_skipped:
            summary += " 未完成项：" + "；".join(safe_skipped[:3])
    return summary + report_links


@tool
def fetch_and_audit_open_supplements(doi: str) -> str:
    """按用户明确要求的 DOI 查找 Europe PMC/NCBI PMC 开放获取补充材料；仅下载可信 PMC OA 清单中的 HTTPS 文件，扫描其中可识别的表格，返回数据线索和下载报告。未找到不代表论文无补充数据。"""
    doi = _normalize_doi(doi)
    if not re.fullmatch(r"10\.\d{4,9}/[^\s<>\[\]]+", doi):
        return "DOI 格式无效；未发起下载。"
    import httpx
    try:
        response = httpx.get("https://www.ebi.ac.uk/europepmc/webservices/rest/search",
                             params={"query": f'DOI:"{doi}"', "format": "json", "resultType": "core"},
                             timeout=10, follow_redirects=True)
        response.raise_for_status()
        records = response.json().get("resultList", {}).get("result", [])
        record = next((r for r in records if _normalize_doi(r.get("doi", "")) == doi), None)
        pmcid = str((record or {}).get("pmcid") or "")
        if not re.fullmatch(r"PMC\d+", pmcid, re.I):
            return ("Europe PMC 未找到该 DOI 对应的开放全文补充材料入口；"
                    "这不代表论文没有补充数据，其他仓储未在本次查询范围内。")
        oa = httpx.get("https://www.ncbi.nlm.nih.gov/pmc/utils/oa/oa.fcgi",
                       params={"id": pmcid}, timeout=10, follow_redirects=True)
        oa.raise_for_status()
        root = ET.fromstring(oa.content)
        package_url = next((node.attrib.get("href", "") for node in root.iter()
                            if node.tag.lower().endswith("link") and
                            node.attrib.get("href", "").lower().endswith((".tar.gz", ".tgz"))), "")
        if not package_url:
            return (f"已找到开放全文记录 {pmcid}，但 OA 清单没有可下载归档；"
                    "未把它记作无补充数据。")
        from urllib.parse import urlsplit
        parsed = urlsplit(package_url)
        if parsed.scheme != "https" or parsed.hostname not in {"ftp.ncbi.nlm.nih.gov", "www.ncbi.nlm.nih.gov"}:
            return "补充材料下载地址不在可信 PMC OA 主机白名单中，已停止下载。"
        with httpx.stream("GET", package_url, timeout=30, follow_redirects=False) as stream:
            stream.raise_for_status()
            size_header = int(stream.headers.get("content-length") or 0)
            if size_header > 100 * 1024 * 1024:
                return "补充材料归档超过 100 MB，本次未下载。"
            chunks = []; total = 0
            for chunk in stream.iter_bytes():
                total += len(chunk)
                if total > 100 * 1024 * 1024:
                    return "补充材料归档超过 100 MB，已停止下载。"
                chunks.append(chunk)
        archive_bytes = b"".join(chunks)
    except Exception as exc:
        return f"公开补充材料获取未完成：{type(exc).__name__}。未把网络失败解释为文献或数据不存在。"
    supported = {".csv", ".tsv", ".xlsx", ".xls", ".xlsm", ".pdf", ".docx"}
    table_findings = []; scanned = []; skipped = []
    try:
        with tarfile.open(fileobj=io.BytesIO(archive_bytes), mode="r:*") as tar:
            for member in tar.getmembers():
                if not member.isfile() or member.size > 30 * 1024 * 1024:
                    continue
                suffix = Path(member.name).suffix.lower()
                if suffix not in supported:
                    continue
                if len(scanned) >= 30 or sum(len(x[1]) for x in scanned) + member.size > 100 * 1024 * 1024:
                    skipped.append(Path(member.name).name + "（扫描预算达到上限）")
                    continue
                handle = tar.extractfile(member)
                if handle is None:
                    continue
                content = handle.read(30 * 1024 * 1024 + 1)
                if len(content) > 30 * 1024 * 1024:
                    skipped.append(Path(member.name).name + "（单文件超过 30 MB）"); continue
                scanned.append((Path(member.name).name, content))
    except Exception as exc:
        return f"已下载 PMC OA 归档，但归档解析失败（{type(exc).__name__}）；没有声明审查通过。"
    for name, content in scanned:
        try:
            suffix = Path(name).suffix.lower()
            if suffix in {".csv", ".tsv"}:
                frame, *_ = _decode_csv(content, 1, [1])
            elif suffix in {".xlsx", ".xlsm"}:
                frame = _pd().read_excel(io.BytesIO(content), sheet_name=0, engine="openpyxl")
            elif suffix == ".xls":
                import python_calamine  # noqa: F401
                frame = _pd().read_excel(io.BytesIO(content), sheet_name=0, engine="calamine")
            elif suffix == ".pdf":
                import pdfplumber
                with pdfplumber.open(io.BytesIO(content)) as pdf:
                    rows = next((t for page in pdf.pages[:50] for t in (page.extract_tables() or []) if t), None)
                if not rows:
                    skipped.append(name + "（未提取到表格）"); continue
                width = max(len(row) for row in rows)
                headers = _normalize_headers(rows, [1], width)
                frame = _pd().DataFrame([(r + [None] * width)[:width] for r in rows[1:100001]], columns=headers)
            else:
                skipped.append(name + "（Word 表格暂未纳入补充材料批处理）"); continue
            findings, _ = _audit_frame(frame, {}, [], [], row_offset=2)
            for finding in findings:
                finding["field"] = f"{name} · {finding.get('field', '')}"
            table_findings.extend(findings)
        except Exception as exc:
            skipped.append(f"{name}（解析失败：{type(exc).__name__}）")
    meta = {"name": f"PMC OA {pmcid}", "sheet": "补充材料", "rows": len(scanned), "columns": 0,
            "scan_errors": skipped}
    report_bytes = _write_audit_xlsx(f"{doi.replace('/', '_')}_补充数据审查.xlsx", meta, table_findings)
    report_url = _save_artifact(f"{doi.replace('/', '_')}_补充数据审查.xlsx", report_bytes,
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    archive_url = _save_artifact(f"{pmcid}_补充材料.tar.gz", archive_bytes, "application/gzip")
    status = f"找到 {len(scanned)} 个支持格式的材料并完成基础扫描，产生 {len(table_findings)} 项复核线索。"
    if not scanned:
        status = "已取得开放归档，但没有识别到可审查的表格材料；不能据此判断没有数据问题。"
    if skipped:
        status += f"另有 {len(skipped)} 项未能审查：" + "；".join(skipped[:10])
    return (f"{status}\n\n来源：Europe PMC / NCBI PMC OA（{pmcid}）。线索不构成错误或不端结论。\n"
            f"[下载原始补充材料归档]({archive_url})\n[下载数据审查报告]({report_url})")


def _clean_reference_value(value: Any) -> str:
    text = str(value or "")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\r\n\t]+", " ", text).strip()[:4000]


def _render_reference_export(rows: list[dict[str, Any]], output_format: str) -> tuple[str, str, str]:
    """Render verified bibliography metadata without inferring missing fields."""
    fmt = str(output_format or "xlsx").strip().lower()
    if fmt == "ris":
        records = []
        for row in rows[:2000]:
            fields = ["TY  - JOUR" if _clean_reference_value(row.get("journal")) else "TY  - GEN"]
            for key, tag in (("title", "TI"), ("journal", "JO"), ("year", "PY"),
                             ("volume", "VL"), ("issue", "IS"), ("pages", "SP"),
                             ("doi", "DO"), ("url", "UR")):
                value = _clean_reference_value(row.get(key, ""))
                if value:
                    fields.append(f"{tag}  - {value}")
            authors = row.get("authors") or []
            if isinstance(authors, str):
                authors = [part.strip() for part in re.split(r"\s*;\s*|\s*\|\s*", authors) if part.strip()]
            if isinstance(authors, list):
                fields.extend(f"AU  - {_clean_reference_value(author)}" for author in authors if _clean_reference_value(author))
            citation = _clean_reference_value(row.get("citation", ""))
            evidence = _clean_reference_value(row.get("evidence", ""))
            claim = _clean_reference_value(row.get("claim", ""))
            notes = " | ".join(x for x in (citation, f"结论: {claim}" if claim else "",
                                             f"证据: {evidence}" if evidence else "") if x)
            if notes:
                fields.append(f"N1  - {notes}")
            fields.append("ER  - ")
            records.append("\r\n".join(fields))
        return "\r\n\r\n".join(records) + "\r\n", ".ris", "application/x-research-info-systems; charset=utf-8"
    if fmt == "bibtex":
        def escape(value: str) -> str:
            value = _clean_reference_value(value)
            substitutions = {"\\": r"\textbackslash{}", "{": r"\{", "}": r"\}",
                             "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
                             "_": r"\_", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
            return "".join(substitutions.get(character, character) for character in value)

        records = []
        for index, row in enumerate(rows[:2000], 1):
            seed = _clean_reference_value(row.get("doi") or row.get("url") or row.get("citation") or f"ref{index}")
            key = "ref_" + hashlib.sha1(seed.encode("utf-8")).hexdigest()[:10]
            fields = []
            for key_name, source_key in (("title", "title"), ("author", "authors"),
                                         ("journal", "journal"), ("year", "year"),
                                         ("volume", "volume"), ("number", "issue"),
                                         ("pages", "pages"), ("doi", "doi"), ("url", "url")):
                value = row.get(source_key, "")
                if isinstance(value, list):
                    value = " and ".join(_clean_reference_value(part) for part in value if _clean_reference_value(part))
                elif source_key == "authors" and isinstance(value, str):
                    value = " and ".join(part.strip() for part in re.split(r"\s*;\s*|\s*\|\s*", value) if part.strip())
                value = escape(value)
                if value:
                    fields.append(f"  {key_name} = {{{value}}}")
            citation = escape(row.get("citation", ""))
            evidence = escape(row.get("evidence", ""))
            claim = escape(row.get("claim", ""))
            note = "; ".join(x for x in (citation, f"Claim: {claim}" if claim else "",
                                              f"Evidence: {evidence}" if evidence else "") if x)
            if note:
                fields.append(f"  note = {{{note}}}")
            entry_type = "article" if _clean_reference_value(row.get("journal")) else "misc"
            records.append("@" + entry_type + "{" + key + ",\n" + ",\n".join(fields) + "\n}")
        return "\n\n".join(records) + "\n", ".bib", "application/x-bibtex; charset=utf-8"
    raise ValueError("format 必须为 xlsx、ris 或 bibtex")


@tool
def export_literature_evidence(title: str, evidence_rows_json: str,
                               output_format: str = "xlsx") -> str:
    """导出已检索/核验的文献证据；output_format 可选 xlsx、ris、bibtex。行可含 citation、title、authors、journal、year、volume、issue、pages、claim、evidence、source、doi、url；缺失字段不补造。"""
    rows = _json_arg(evidence_rows_json, list, [])
    rows = [row for row in rows[:2000] if isinstance(row, dict)]
    output_format = str(output_format or "xlsx").strip().lower()
    if output_format in {"ris", "bibtex"}:
        body, extension, mime = _render_reference_export(rows, output_format)
        filename = _safe_name(title, "文献证据") + extension
        url = _save_artifact(filename, body.encode("utf-8"), mime)
        return f"文献引用已导出（{len(rows)} 条）。[下载 {output_format.upper()}]({url})"
    if output_format != "xlsx":
        return "[ERR_PARSE] output_format 只支持 xlsx、ris 或 bibtex。"
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "文献证据"
    headers = ["文献", "支持的结论", "原文/证据摘要", "来源", "DOI", "链接"]
    ws.append([_safe_spreadsheet_value(value) for value in headers])
    for item in rows:
        ws.append([_safe_spreadsheet_value(item.get(key, "")) for key in
                   ("citation", "claim", "evidence", "source", "doi", "url")])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEF", (34, 42, 70, 20, 25, 55)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO(); wb.save(out)
    filename = _safe_name(title, "文献证据") + ".xlsx"
    url = _save_artifact(filename, out.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return f"文献证据表已生成（{len(rows)} 条）。[下载 XLSX]({url})"


def _export_markdown(title: str, content: str, kind: str, svg_code: str = "") -> str:
    body = str(content or "").strip()
    svg_source = str(svg_code or "").strip()
    if not svg_source:
        svg_block = re.search(r"```(?:svg|xml)\s*\n([\s\S]*?)```", body, re.I)
        if svg_block:
            svg_source = svg_block.group(1).strip()
        else:
            inline_svg = re.search(r"<svg\b[\s\S]*?</svg\s*>", body, re.I)
            if inline_svg:
                svg_source = inline_svg.group(0)
    if svg_source:
        fenced = re.fullmatch(r"```(?:svg|xml)?\s*([\s\S]*?)\s*```", svg_source, re.I)
        if fenced:
            svg_source = fenced.group(1).strip()
        else:
            start = svg_source.lower().find("<svg")
            end = svg_source.lower().rfind("</svg>")
            if start >= 0 and end >= start:
                svg_source = svg_source[start:end + len("</svg>")]
        svg_code = _sanitize_svg_for_export(svg_source)
        if not svg_code:
            return "[SVG_FORMAT_ERROR] SVG 不完整或未通过安全校验，未生成图形文件。"
    if not body and svg_code:
        body = f"## 图示\n\n```svg\n{svg_code}\n```"
    if not body:
        return "[ERR_PARSE] 导出内容不能为空；请提供正文或有效 SVG。"
    supplied_title = str(title or "").strip()
    heading = re.search(r"^#\s+(.+)$", body, re.M)
    generic = {"汇总", "导出文件", "课题记录", "文件", "结果", "报告"}
    if (not supplied_title or supplied_title in generic) and heading:
        supplied_title = heading.group(1).strip()
    if not supplied_title or supplied_title in generic:
        svg_title = re.search(r"<title\b[^>]*>([\s\S]*?)</title>", svg_code, re.I)
        if svg_title:
            supplied_title = html.unescape(re.sub(r"<[^>]+>", "", svg_title.group(1))).strip()
    final_title = supplied_title or kind
    if svg_code and "<svg" not in body.lower():
        body += "\n\n## 配图\n\n" + str(svg_code).strip()
    rendered = _markdown_to_safe_html(body)
    document = f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(final_title)}</title><style>
body{{max-width:980px;margin:40px auto;padding:0 24px;color:#242424;font:16px/1.75 system-ui,"Microsoft YaHei",sans-serif}}
h1,h2,h3{{line-height:1.35;margin-top:1.6em}} table{{border-collapse:collapse;width:100%;margin:1em 0;font-size:14px}}
th,td{{border:1px solid #ddd;padding:8px 10px;text-align:left;vertical-align:top}} th{{background:#f6f6f6}}
pre{{overflow:auto;background:#f6f6f6;padding:14px;border-radius:8px}} code{{background:#f6f6f6;padding:2px 4px;border-radius:4px}}
a{{color:#d97706}} .figure{{max-width:100%;overflow:auto;margin:20px 0}} .figure svg{{display:block;max-width:100%;height:auto}}
.muted{{color:#777;font-size:13px}}
</style></head><body><h1>{html.escape(final_title)}</h1>{rendered}</body></html>'''
    filename = _safe_name(final_title, kind) + ".html"
    url = _save_artifact(filename, document.encode("utf-8"), "text/html; charset=utf-8")
    links = [f"[下载 HTML]({url})"]
    if svg_code:
        svg_filename = _safe_name(final_title, kind) + ".svg"
        svg_url = _save_artifact(svg_filename, svg_code.encode("utf-8"), "image/svg+xml; charset=utf-8")
        links.append(f"[下载 SVG]({svg_url})")
    suffix = "（SVG 已内嵌）" if svg_code else ""
    return f"{kind}已生成 HTML{suffix}：" + " ｜ ".join(links)


@tool
def export_paper_analysis(analysis_markdown: str = "", title: str = "", svg_code: str = "") -> str:
    """将本轮论文解读或绘图导出为 HTML；title 可省略并根据正文标题或 SVG 标题自动命名，svg_code 可单独用于导出绘图。"""
    return _export_markdown(title, analysis_markdown, "论文解读", svg_code)


@tool
def export_experiment_design(design_markdown: str = "", title: str = "", svg_code: str = "") -> str:
    """导出实验设计 HTML；注明待确认假设。适合随文呈现的 SVG 可放入 svg_code 自动内嵌。"""
    return _export_markdown(title, design_markdown, "实验设计草案", svg_code)


@tool
def save_research_note(project_name: str, note: str) -> str:
    """仅在用户明确要求时，向该账号跨会话共享的课题记录写入一条备注。"""
    sid = str(get_session_id() or "")
    owner = get_user_id()
    project = str(project_name or "").strip()[:120]
    text = str(note or "").strip()[:10000]
    if not sid or not owner or not project or not text:
        return "[ERR_PARSE] 需要有效登录账号、会话、课题名称和非空记录"
    with sqlite3.connect(str(_db_path())) as conn:
        conn.execute("INSERT INTO research_notes(session_id,owner_user_id,project_name,note,created_at) "
                     "VALUES(?,?,?,?,?)", (sid, int(owner), project, text, time.time()))
    return f"已将记录保存到账号内课题“{project}”，可在该账号的其他会话回顾；不会写入共享知识库。"


@tool
def recall_research_notes(project_name: str, limit: int = 20) -> str:
    """回顾当前登录账号在不同会话中已明确保存的指定课题记录。"""
    owner = get_user_id()
    project = str(project_name or "").strip()[:120]
    if not owner or not project:
        return "[ERR_PARSE] 需要有效登录账号和课题名称"
    with sqlite3.connect(str(_db_path())) as conn:
        rows = conn.execute("SELECT note,created_at FROM research_notes WHERE owner_user_id=? "
                            "AND project_name=? ORDER BY id DESC LIMIT ?",
                            (int(owner), project, max(1, min(int(limit), 50)))).fetchall()
    if not rows:
        return f"当前账号没有课题“{project}”的已保存记录。"
    return "\n".join(f"- [{datetime.fromtimestamp(ts, timezone.utc).astimezone().isoformat(timespec='minutes')}] {note}"
                     for note, ts in reversed(rows))


@tool
def export_research_notes(project_name: str, limit: int = 100, svg_code: str = "") -> str:
    """将课题记录导出为按课题命名的 HTML 文件；若记录配有 SVG，将代码放入 svg_code 自动安全清理并内嵌。"""
    owner = get_user_id()
    project = str(project_name or "").strip()[:120]
    if not owner or not project:
        return "[ERR_PARSE] 需要有效登录账号和课题名称"
    with sqlite3.connect(str(_db_path())) as conn:
        rows = conn.execute("SELECT note,created_at FROM research_notes WHERE owner_user_id=? "
                            "AND project_name=? ORDER BY id ASC LIMIT ?",
                            (int(owner), project, max(1, min(int(limit), 500)))).fetchall()
    if not rows:
        return f"当前账号没有课题“{project}”的记录可导出。"
    body = "\n\n".join(
        f"## {datetime.fromtimestamp(ts, timezone.utc).astimezone().isoformat(timespec='minutes')}\n\n{note}"
        for note, ts in rows)
    return _export_markdown(project + "课题记录", body, "课题记录", svg_code)


@tool
def export_corpus_gap_report(title: str, gaps_json: str) -> str:
    """导出本轮已观察到的本地语料缺口或网页获取失败记录。每项可含 topic、doi_or_title、failure、local_result、source_url；只记录实际观察到的信息。"""
    gaps = _json_arg(gaps_json, list, [])
    fields = ["topic", "doi_or_title", "failure", "local_result", "source_url"]
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for item in gaps[:5000]:
        row = item if isinstance(item, dict) else {}
        writer.writerow({field: _safe_spreadsheet_value(row.get(field, "")) for field in fields})
    url = _save_artifact(_safe_name(title, "语料缺口") + ".csv",
                         output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8")
    return f"语料缺口报告已生成（{min(len(gaps), 5000)} 条）。[下载 CSV]({url})"

